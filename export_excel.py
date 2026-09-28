"""
export_excel.py
Builds a multi-sheet Excel workbook from the analysis outputs:
  - Raw data summary tables
  - Pivot-style summary tables
  - Native Excel charts
  - A "Key Findings" sheet
"""

import pandas as pd
import os
import pyodbc
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
pyodbc.pooling = False
conn = pyodbc.connect(
    r"DRIVER={ODBC Driver 18 for SQL Server};"
    r"SERVER=.\SQLEXPRESS;"
    r"DATABASE=HealthCare;"
    r"Trusted_Connection=yes;"
    r"Encrypt=yes;"
    r"TrustServerCertificate=yes;"
    r"MARS_Connection=no;"
    r"APP=CarePath Excel Report;",
    timeout=5,
)


def read_sql(query):
    cursor = conn.cursor()
    cursor.execute(query)
    columns = [column[0] for column in cursor.description]
    return pd.DataFrame.from_records(cursor.fetchall(), columns=columns)


monthly_volume = read_sql("""
    SELECT CONVERT(char(7), appointment_date, 126) AS month,
           COUNT(*) AS num_appointments
    FROM appointments
    WHERE appointment_date IS NOT NULL
    GROUP BY CONVERT(char(7), appointment_date, 126)
    ORDER BY month
""")

monthly_revenue = read_sql("""
    SELECT CONVERT(char(7), billing_date, 126) AS month,
           SUM(amount) AS amount
    FROM billing
    WHERE billing_date IS NOT NULL
    GROUP BY CONVERT(char(7), billing_date, 126)
    ORDER BY month
""")

patients = read_sql("SELECT dob, city FROM patients WHERE dob IS NOT NULL")
patients["dob"] = pd.to_datetime(patients["dob"])
patients["age"] = ((pd.Timestamp.today().normalize() - patients["dob"]).dt.days / 365.25).astype(int)
patients["age_group"] = pd.cut(
    patients["age"],
    bins=[0, 18, 40, 60, 120],
    labels=["Under 18", "18-39", "40-59", "60+"],
    right=False,
)
age_seg = patients["age_group"].value_counts().rename_axis("age_group").reset_index(name="num_patients")
city_seg = patients["city"].value_counts().rename_axis("city").reset_index(name="num_patients")

dept_revenue = read_sql("""
    SELECT dp.name AS department, SUM(b.amount) AS revenue
    FROM billing b
    JOIN diagnoses dg ON b.diagnosis_id = dg.diagnosis_id
    JOIN doctors d ON dg.doctor_id = d.doctor_id
    JOIN departments dp ON d.department_id = dp.department_id
    GROUP BY dp.name
    ORDER BY revenue DESC
""")

# Extra pivot: appointment status breakdown by department
q = """
SELECT dp.name AS department, a.status, COUNT(*) AS n
FROM appointments a
JOIN doctors d ON a.doctor_id = d.doctor_id
JOIN departments dp ON d.department_id = dp.department_id
GROUP BY dp.name, a.status
"""
status_by_dept = read_sql(q)
status_pivot = status_by_dept.pivot_table(index="department", columns="status", values="n", fill_value=0)

no_show_summary = read_sql("""
    SELECT COUNT(*) AS total_appointments,
           SUM(CASE
                 WHEN LOWER(REPLACE(REPLACE(status, ' ', ''), '-', '')) = 'noshow' THEN 1
                 ELSE 0
               END) AS no_show_appointments
    FROM appointments
""").iloc[0]
total_appointments = int(no_show_summary["total_appointments"])
no_show_appointments = int(no_show_summary["no_show_appointments"] or 0)
no_show_rate = 100 * no_show_appointments / total_appointments if total_appointments else 0

wb = Workbook()
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

def write_df(ws, df, start_row=1, start_col=1, with_index=False):
    rows = dataframe_to_rows(df, index=with_index, header=True)
    r = start_row
    for row in rows:
        if row == [None] and with_index:
            continue
        for c, val in enumerate(row, start=start_col):
            ws.cell(row=r, column=c, value=val)
        r += 1
    # header styling
    for c in range(start_col, start_col + len(df.columns) + (1 if with_index else 0)):
        cell = ws.cell(row=start_row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    return r  # next free row

# ---------------- Sheet 1: Key Findings ----------------
ws = wb.active
ws.title = "Key Findings"
ws.column_dimensions["A"].width = 90
findings = [
    "CLINIC OPERATIONS — KEY FINDINGS",
    "",
    f"1. Peak appointment month: {monthly_volume.loc[monthly_volume['num_appointments'].idxmax(), 'month']}",
    f"2. Peak revenue month: {monthly_revenue.loc[monthly_revenue['amount'].idxmax(), 'month']} "
    f"(EGP {monthly_revenue['amount'].max():,.2f})",
    f"3. Highest-revenue department: {dept_revenue.iloc[0]['department']} "
    f"(EGP {dept_revenue.iloc[0]['revenue']:,.2f})",
    f"4. Largest patient age segment: {age_seg.iloc[0]['age_group']} ({age_seg.iloc[0]['num_patients']} patients)",
    f"5. Most represented city: {city_seg.iloc[0]['city']} ({city_seg.iloc[0]['num_patients']} patients)",
    f"6. No-show rate: {no_show_rate:.1f}% ({no_show_appointments} of {total_appointments} appointments).",
    "7. See the 'Appointment Status' sheet for the appointment-status breakdown by department.",
]
for i, line in enumerate(findings, start=1):
    cell = ws.cell(row=i, column=1, value=line)
    if i == 1:
        cell.font = Font(bold=True, size=14)
    cell.alignment = Alignment(wrap_text=True)

# ---------------- Sheet 2: Monthly Trends (+ line chart) ----------------
ws2 = wb.create_sheet("Monthly Trends")
next_row = write_df(ws2, monthly_volume)
write_df(ws2, monthly_revenue, start_col=4)

chart1 = LineChart()
chart1.title = "Monthly Appointment Volume"
chart1.y_axis.title = "Appointments"
chart1.x_axis.title = "Month"
data = Reference(ws2, min_col=2, min_row=1, max_row=len(monthly_volume) + 1)
cats = Reference(ws2, min_col=1, min_row=2, max_row=len(monthly_volume) + 1)
chart1.add_data(data, titles_from_data=True)
chart1.set_categories(cats)
ws2.add_chart(chart1, "H2")

chart2 = LineChart()
chart2.title = "Monthly Revenue (EGP)"
data2 = Reference(ws2, min_col=5, min_row=1, max_row=len(monthly_revenue) + 1)
cats2 = Reference(ws2, min_col=4, min_row=2, max_row=len(monthly_revenue) + 1)
chart2.add_data(data2, titles_from_data=True)
chart2.set_categories(cats2)
ws2.add_chart(chart2, "H18")

# ---------------- Sheet 3: Department Revenue (+ bar chart) ----------------
ws3 = wb.create_sheet("Department Revenue")
write_df(ws3, dept_revenue)
chart3 = BarChart()
chart3.title = "Revenue by Department"
chart3.y_axis.title = "Revenue (EGP)"
data3 = Reference(ws3, min_col=2, min_row=1, max_row=len(dept_revenue) + 1)
cats3 = Reference(ws3, min_col=1, min_row=2, max_row=len(dept_revenue) + 1)
chart3.add_data(data3, titles_from_data=True)
chart3.set_categories(cats3)
ws3.add_chart(chart3, "D2")

# ---------------- Sheet 4: Patient Segmentation (+ pie chart) ----------------
ws4 = wb.create_sheet("Patient Segmentation")
write_df(ws4, age_seg)
write_df(ws4, city_seg, start_col=4)

chart4 = PieChart()
chart4.title = "Patients by Age Group"
data4 = Reference(ws4, min_col=2, min_row=1, max_row=len(age_seg) + 1)
cats4 = Reference(ws4, min_col=1, min_row=2, max_row=len(age_seg) + 1)
chart4.add_data(data4, titles_from_data=True)
chart4.set_categories(cats4)
ws4.add_chart(chart4, "H2")

chart5 = BarChart()
chart5.title = "Patients by City"
data5 = Reference(ws4, min_col=5, min_row=1, max_row=len(city_seg) + 1)
cats5 = Reference(ws4, min_col=4, min_row=2, max_row=len(city_seg) + 1)
chart5.add_data(data5, titles_from_data=True)
chart5.set_categories(cats5)
ws4.add_chart(chart5, "H18")

# ---------------- Sheet 5: Appointment Status Pivot (+ chart) ----------------
ws5 = wb.create_sheet("Appointment Status")
status_pivot_reset = status_pivot.reset_index()
write_df(ws5, status_pivot_reset)

chart6 = BarChart()
chart6.type = "col"
chart6.grouping = "stacked"
chart6.overlap = 100
chart6.title = "Appointment Status by Department"
n_cols = len(status_pivot_reset.columns)
data6 = Reference(ws5, min_col=2, max_col=n_cols, min_row=1, max_row=len(status_pivot_reset) + 1)
cats6 = Reference(ws5, min_col=1, min_row=2, max_row=len(status_pivot_reset) + 1)
chart6.add_data(data6, titles_from_data=True)
chart6.set_categories(cats6)
ws5.add_chart(chart6, "H2")

report_path = os.environ.get("CLINIC_REPORT_PATH", os.path.join(BASE_DIR, "clinic_report.xlsx"))
wb.save(report_path)
print("Excel report saved:", report_path)
conn.close()
